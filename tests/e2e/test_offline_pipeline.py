"""E2E без модели и GPU: шаблон + контент -> 3 варианта; на трёх шаблонах датасета и на незнакомом шаблоне."""
import time

import pytest
from pptx import Presentation

from deckgen.models import Severity

HARD = {"integrity.file_opens", "integrity.raster_slide", "integrity.leftover_placeholder", "template.layout",
        "template.font_family", "layout.out_of_bounds", "content.numbers_grounded"}


def _run(settings, template, pack):
    from deckgen.pipeline.orchestrator import Pipeline
    t0 = time.perf_counter()
    res = Pipeline(settings).run(template, pack)
    return res, time.perf_counter() - t0


@pytest.mark.parametrize("idx", [0, 1, 2])
def test_dataset_template(settings, templates, pack, idx):
    if idx >= len(templates):
        pytest.skip()
    res, dt = _run(settings, templates[idx], pack)
    assert dt < settings.run.max_generation_seconds
    assert set(res.variants) == {"balanced", "visual", "compact"}
    sizes = {len(v.plan.slides) for v in res.variants.values()}
    picks = {tuple(s.template_slide for s in v.plan.slides) for v in res.variants.values()}
    assert len(picks) == 3                                          # варианты визуально различаются (другие макеты)
    assert len(sizes) >= 2                                          # и составом
    for v in res.variants.values():
        prs = Presentation(v.files["pptx"])
        assert 8 <= len(prs.slides) <= 16
        errors = [f for f in v.audit.findings if f.severity == Severity.error]
        assert not [f for f in errors if f.check_id in HARD], [(f.slide_index, f.check_id, f.message) for f in errors]
        asses = {f.check_id for f in v.audit.findings}
        assert "integrity.raster_slide" not in asses
        assert v.files["html"].endswith(".html")


def test_unknown_blank_template(settings, blank_template, pack):
    res, _ = _run(settings, blank_template, pack)
    for v in res.variants.values():
        assert len(Presentation(v.files["pptx"]).slides) >= 8
        assert not [f for f in v.audit.findings if f.check_id in HARD and f.severity == Severity.error]


def test_free_text_input(settings, templates):
    from deckgen.content.pack import pack_from_text
    text = ("Предлагаем внедрить единый портал заявок для HR. Сейчас заявка обрабатывается 5 дней, 40% заявок теряются. "
            "Пилот в двух филиалах сократил срок до 2 дней. Стоимость внедрения — 3 млн рублей, окупаемость 8 месяцев. "
            "Просим одобрить запуск во всех 12 филиалах в первом квартале.")
    pack = pack_from_text(text, audience="правление", target_slides=10)
    res, _ = _run(settings, templates[0], pack)
    assert all(len(v.plan.slides) >= 6 for v in res.variants.values())
