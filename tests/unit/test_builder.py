"""Сборка: клон образца сохраняет дизайн, стиль рана-образца, нативные график/таблица, чистка макетов."""
from pathlib import Path

from pptx import Presentation

from deckgen.models import (ChartSeries, ChartSpec, DeckPlan, PlannedSlide, TableSpec, TextFill, VisualFill)
from deckgen.ooxml.reader import PresentationReader
from deckgen.render.pptx_builder import build_pptx
from tests.conftest import by_name


def _plan(tm, slides):
    return DeckPlan(deck_id="t", variant="balanced", template_id=tm.design.template_id, template_file=tm.design.source_file,
                    content_pack_id="p", slides=slides)


def test_clone_keeps_design_and_run_style(parsed, tmp_path):
    tm = by_name(parsed, "WorkSpace")
    t = tm.slides[5]                                              # 3 карточки: подложки 212121, заголовки 0077FF
    heading = next(s for s in t.texts if s.role.value == "heading")
    body = next(s for s in t.texts if s.role.value == "body")
    ps = PlannedSlide(index=0, template_slide=5, layout_part=t.layout_part, kind=t.kind,
                      texts=[TextFill(shape_id=heading.shape_id, paragraphs=["Три клиента"], role=heading.role),
                             TextFill(shape_id=body.shape_id, paragraphs=["Пилот у ритейла, банка и промышленности"], role=body.role)])
    out = build_pptx(_plan(tm, [ps]), tm, tmp_path / "o.pptx")
    src = PresentationReader(tm.design.source_file).read_slide(5)
    dst = PresentationReader(out).read_slide(0)
    assert len(dst.shapes) == len(src.shapes)                     # вся графика образца на месте
    h = dst.by_id()[heading.shape_id]
    assert h.text == "Три клиента" and h.main_run.color == "0077FF" and h.main_run.size == heading.size_pt
    assert dst.by_id()[heading.shape_id].fill == src.by_id()[heading.shape_id].fill
    assert len(Presentation(str(out)).slides) == 1                # образцы шаблона удалены
    assert out.stat().st_size < Path(tm.design.source_file).stat().st_size / 2   # неиспользованные макеты вычищены


def test_empty_run_is_not_used_as_style(parsed, tmp_path):
    tm = by_name(parsed, "WorkSpace")
    t = tm.slides[10]
    slot = next(s for s in t.texts if s.shape_id == "566")        # в образце: белый текст, пустые раны — чёрные
    ps = PlannedSlide(index=0, template_slide=10, layout_part=t.layout_part, kind=t.kind,
                      texts=[TextFill(shape_id="566", paragraphs=["Резюме встреч"], role=slot.role)])
    out = build_pptx(_plan(tm, [ps]), tm, tmp_path / "o.pptx")
    assert PresentationReader(out).read_slide(0).by_id()["566"].main_run.color == "FFFFFF"


def test_native_chart_and_template_table(parsed, tmp_path):
    tm = by_name(parsed, "WorkSpace")
    tt, ct = tm.slides[13], tm.slides[19]
    table = PlannedSlide(index=0, template_slide=13, layout_part=tt.layout_part, kind=tt.kind,
                         visual=VisualFill(kind="table", box=tt.visual.box, table_shape_id=tt.visual.table_shape_id,
                                           table=TableSpec(columns=["Клиент", "DAU, %"], rows=[["Ритейл", "71"], ["Банк", "64"]])))
    chart = PlannedSlide(index=1, template_slide=19, layout_part=ct.layout_part, kind=ct.kind,
                         visual=VisualFill(kind="chart", box=ct.visual.box, remove_ids=ct.visual.replace_ids,
                                           chart=ChartSpec(chart_type="line", categories=["Май", "Июнь"],
                                                           series=[ChartSeries(name="Пользователи", values=[120, 340], color="0077FF")],
                                                           x_title="Месяц", y_title="чел."), text_color="FFFFFF"))
    out = build_pptx(_plan(tm, [table, chart]), tm, tmp_path / "o.pptx")
    prs = Presentation(str(out))
    tbl = next(s for s in prs.slides[0].shapes if s.has_table).table
    assert len(tbl.rows) == 3 and len(tbl.columns) == 2 and tbl.cell(1, 0).text == "Ритейл"
    ch = next(s for s in prs.slides[1].shapes if s.has_chart).chart
    assert list(ch.plots[0].categories) == ["Май", "Июнь"]
    assert not any(s.shape_type == 13 and str(s.shape_id) in ct.visual.replace_ids for s in prs.slides[1].shapes)


def test_item_delete_and_reflow(parsed, tmp_path, settings, pack):
    from deckgen.composing.composer import _reflow
    tm = by_name(parsed, "WorkSpace")
    t = tm.slides[5]
    moves = _reflow(t, 2)                                          # 3 карточки -> 2, по центру прежнего ряда
    assert moves
    xs = sorted(b.x for b in moves.values())
    assert xs[0] > t.items[0].box.x
