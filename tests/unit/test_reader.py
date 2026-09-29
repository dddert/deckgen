"""OOXML-ридер: трансформации групп, наследование плейсхолдеров, fontScale, альфа заливки, цвет из стиля."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

from deckgen.ooxml.colors import ColorContext, Theme, contrast_ratio
from deckgen.ooxml.reader import PresentationReader, q


def test_group_child_transform(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    grp = slide.shapes.add_group_shape()
    child = grp.shapes.add_textbox(Emu(0), Emu(0), Emu(914400), Emu(914400))
    child.text_frame.text = "x"
    xfrm = grp._element.find(q("p:grpSpPr")).find(q("a:xfrm"))
    # группа на слайде 1×1 in в точке (1 in, 1 in), а дети описаны в пространстве 2×2 in -> масштаб 0.5
    for tag, attrs in (("a:off", {"x": "914400", "y": "914400"}), ("a:ext", {"cx": "914400", "cy": "914400"}),
                       ("a:chOff", {"x": "0", "y": "0"}), ("a:chExt", {"cx": "1828800", "cy": "1828800"})):
        el = xfrm.find(q(tag))
        for k, v in attrs.items():
            el.set(k, v)
    p = tmp_path / "g.pptx"
    prs.save(str(p))
    r = PresentationReader(p)
    s = next(x for x in r.read_slide(0).shapes if x.kind == "sp")
    W, H = r.slide_w_in, r.slide_h_in
    assert abs(s.box.x * W - 1.0) < 1e-3 and abs(s.box.w * W - 0.5) < 1e-3
    assert abs(s.box.y * H - 1.0) < 1e-3 and abs(s.box.h * H - 0.5) < 1e-3
    assert r.read_slide(0).shapes[0].text == ""            # у группы нет «своего» текста (v1 дублировал детей)


def test_placeholder_inherits_geometry_and_style(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])      # Title and Content: у плейсхолдеров нет своего xfrm
    slide.shapes.title.text = "Заголовок"
    p = tmp_path / "ph.pptx"
    prs.save(str(p))
    shapes = PresentationReader(p).read_slide(0).shapes
    title = next(s for s in shapes if s.ph_type == "title")
    assert title.box.w > 0.5 and title.box.h > 0.05          # геометрия пришла из макета
    assert title.main_run.size == 44.0                        # кегль — из titleStyle мастера
    assert title.main_run.family == "Calibri"                 # +mj-lt -> major font темы


def test_font_scale_and_alpha(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(3000000), Emu(900000))
    tb.text_frame.text = "43%"
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(40)
    body = tb.text_frame._txBody.find(q("a:bodyPr"))
    for c in list(body):
        body.remove(c)
    na = body.makeelement(q("a:normAutofit"), {"fontScale": "50000"})
    body.append(na)
    tb.fill.solid()
    tb.fill.fore_color.rgb = RGBColor(0, 0x77, 0xFF)
    clr = tb._element.find(q("p:spPr")).find(q("a:solidFill"))[0]
    clr.append(clr.makeelement(q("a:alpha"), {"val": "20000"}))
    p = tmp_path / "fs.pptx"
    prs.save(str(p))
    s = next(x for x in PresentationReader(p).read_slide(0).shapes if x.text)
    assert s.main_run.size == 20.0                            # 40 pt × fontScale 50%
    assert abs(s.fill_alpha - 0.2) < 1e-6


def test_scheme_colors_and_contrast():
    ctx = ColorContext(Theme(colors={"dk1": "000000", "lt1": "FFFFFF", "accent1": "0077FF"}))
    assert ctx.scheme("tx1") == "000000" and ctx.scheme("bg1") == "FFFFFF"
    assert round(contrast_ratio("FFFFFF", "000000"), 1) == 21.0
