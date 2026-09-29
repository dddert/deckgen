"""Контент-пакет: из папки (формат v1) или из свободного текста пользователя + файлов.

Формат папки (совместим с v1):
  brief.md     — YAML front-matter (id, title, purpose, audience, language, target_slides, tone, must_include) + текст
  facts.yaml   — [{id, text, value?, unit?, source?}]
  data/*.csv|*.xlsx — таблицы (первая строка — заголовки; имя файла = id)
  assets/*     — картинки (logo_*, icon_*, screenshot_* — роль по имени)

Свободный текст: факты (предложения с цифрами) извлекаются детерминированно — это источник для
сверки цифр; таблицы берутся из приложенных csv/xlsx, остальное структурирует модель в skeleton.
"""
from __future__ import annotations

import csv
import hashlib
import re
from pathlib import Path

import yaml

from ..models import Asset, ContentPack, DataTable, DeckPurpose, Fact
from .grounding import numbers_in

_IMG = (".png", ".jpg", ".jpeg", ".webp")
_TXT = (".md", ".txt")
_PURPOSE_WORDS = {
    DeckPurpose.feature: r"фич|функци|feature",
    DeckPurpose.product: r"продукт|запуск|релиз|product",
    DeckPurpose.initiative: r"инициатив|предлага|initiative",
    DeckPurpose.report: r"отч[её]т|итоги|результаты квартал|report",
}


def load_pack(folder: str | Path) -> ContentPack:
    folder = Path(folder)
    meta, brief = _front_matter((folder / "brief.md").read_text(encoding="utf-8")) if (folder / "brief.md").exists() else ({}, "")
    facts = []
    if (folder / "facts.yaml").exists():
        facts = [Fact(**f) for f in (yaml.safe_load((folder / "facts.yaml").read_text(encoding="utf-8")) or [])]
    tables = []
    if (folder / "data").exists():
        for p in sorted((folder / "data").iterdir()):
            tables += read_tables(p)
    assets = _assets(folder / "assets")
    extra = "\n".join(p.read_text(encoding="utf-8") for p in sorted(folder.glob("*.txt")))
    source = "\n".join([brief, extra] + [f.text for f in facts])
    return ContentPack(
        id=str(meta.get("id", folder.name)), title=str(meta.get("title", folder.name)),
        purpose=_purpose(meta.get("purpose"), brief), audience=str(meta.get("audience", "")),
        brief=brief.strip(), language=str(meta.get("language", "ru")), target_slides=meta.get("target_slides"),
        tone=meta.get("tone"), must_include=list(meta.get("must_include", []) or []),
        facts=facts or extract_facts(source), tables=tables, assets=assets, source_text=source,
    )


def pack_from_text(text: str, *, title: str | None = None, purpose: str | None = None, audience: str = "",
                   target_slides: int | None = None, files: list[Path] | None = None, language: str = "ru") -> ContentPack:
    files = files or []
    chunks, tables, assets = [text.strip()], [], []
    for p in files:
        suf = p.suffix.lower()
        if suf in (".csv", ".xlsx", ".xlsm"):
            tables += read_tables(p)
        elif suf in _IMG:
            assets.append(_asset(p))
        elif suf in _TXT:
            chunks.append(p.read_text(encoding="utf-8", errors="ignore"))
        elif suf == ".docx":
            chunks.append(_docx_text(p))
        elif suf == ".pdf":
            chunks.append(_pdf_text(p))
    source = "\n\n".join(c for c in chunks if c.strip())
    first = next((ln.strip(" #*") for ln in source.splitlines() if ln.strip()), "Презентация")
    if not title:
        from ..composing.fitter import truncate_words
        title = first if len(first) <= 90 else truncate_words(first, 88)
    pid = "pack_" + hashlib.md5(source.encode()).hexdigest()[:8]
    return ContentPack(id=pid, title=title, purpose=_purpose(purpose, source), audience=audience, brief=source,
                       language=language, target_slides=target_slides, facts=extract_facts(source),
                       tables=tables, assets=assets, source_text=source)


def extract_facts(text: str, limit: int = 40) -> list[Fact]:
    """Каждое предложение с цифрой — факт. Детерминированно: модель цифры не придумывает, а цитирует."""
    from .textfix import unwrap
    sents = re.split(r"(?<=[.!?…])\s+|\n+", unwrap(text))   # фраза, перенесённая на две строки, — один факт
    out, seen = [], set()
    for s in sents:
        s = s.strip(" -•*\t")
        if len(s) < 12 or not numbers_in(s) or s.lower() in seen:
            continue
        if re.fullmatch(r"[\d\s.,:;%()-]+", s):
            continue
        seen.add(s.lower())
        m = re.search(r"(\d+(?:[.,]\d+)?)\s*(%|процент\w*|ч(?:ас\w*)?\b|мин\w*|сек\w*|млн|млрд|тыс\w*|₽|руб\w*|\$|дн\w*|месяц\w*|раз\w*)?", s)
        out.append(Fact(id=f"f{len(out) + 1}", text=s[:300], value=m.group(1) if m else None,
                        unit=(m.group(2) if m and m.group(2) else None)))
        if len(out) >= limit:
            break
    return out


def read_tables(p: Path) -> list[DataTable]:
    suf = p.suffix.lower()
    if suf == ".csv":
        with p.open(encoding="utf-8-sig", newline="") as f:
            sample = f.read(2048)
            f.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            rows = [r for r in csv.reader(f, dialect) if any(c.strip() for c in r)]
        return [_table(p.stem, rows)] if len(rows) >= 2 else []
    if suf in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(p, read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            rows = [[("" if v is None else str(v)) for v in r] for r in ws.iter_rows(values_only=True)]
            rows = [r for r in rows if any(c.strip() for c in r)]
            if len(rows) >= 2:
                out.append(_table(f"{p.stem}_{ws.title}" if len(wb.worksheets) > 1 else p.stem, rows))
        return out
    return []


def _table(tid: str, rows: list[list[str]]) -> DataTable:
    tid = re.sub(r"[^\w]+", "_", tid).strip("_").lower() or "table"
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    columns = [c.strip() for c in rows[0]]
    return DataTable(id=tid, title=_table_name(tid, columns), columns=columns,
                     rows=[[c.strip() for c in r] for r in rows[1:]])


def _table_name(tid: str, columns: list[str]) -> str:
    """Имя файла («adoption_by_month») на русском слайде — брак: берём названия столбцов с показателями."""
    stem = tid.replace("_", " ")
    if re.search(r"[а-яё]", stem, re.I) or not any(re.search(r"[а-яё]", c, re.I) for c in columns):
        return stem
    name = " и ".join([c for c in columns[1:] if c][:2] or columns[:1])
    return name[:1].upper() + name[1:]


def _front_matter(text: str) -> tuple[dict, str]:
    if text.startswith("---"):
        _, fm, body = text.split("---", 2)
        return yaml.safe_load(fm) or {}, body
    return {}, text


def _purpose(value, text: str) -> DeckPurpose:
    if value:
        try:
            return DeckPurpose(str(value))
        except ValueError:
            pass
    for p, rx in _PURPOSE_WORDS.items():
        if re.search(rx, text[:2000], re.I):
            return p
    return DeckPurpose.project


def _asset(p: Path) -> Asset:
    stem = p.stem.lower()
    kind = "logo" if "logo" in stem else "icon" if "icon" in stem else "screenshot" if ("screen" in stem or "скрин" in stem) else "photo"
    return Asset(id=hashlib.md5(p.read_bytes()).hexdigest()[:10], kind=kind, path=str(p), tags=[p.stem])


def _assets(folder: Path) -> list[Asset]:
    if not folder.exists():
        return []
    return [_asset(p) for p in sorted(folder.iterdir()) if p.suffix.lower() in _IMG]


def _docx_text(p: Path) -> str:
    try:
        import docx  # python-docx, extra [docs]
        return "\n".join(par.text for par in docx.Document(str(p)).paragraphs)
    except Exception:  # noqa: BLE001
        return ""


def _pdf_text(p: Path) -> str:
    try:
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(str(p))
        return "\n".join(page.get_textpage().get_text_range() for page in doc)
    except Exception:  # noqa: BLE001
        return ""
