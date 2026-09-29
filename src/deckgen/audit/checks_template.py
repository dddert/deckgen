"""Шаблон (Приложение 1): гарнитуры, кегли, палитра, макет, логотип/колонтитул, контраст (APCA).
То, что в точности совпадает со слайдом-образцом, помечается «как в шаблоне» (замысел дизайнера)."""
from __future__ import annotations

from ..models import CheckCategory, Finding, PictureRole, Severity
from ..ooxml.colors import apca_lc, contrast_ratio, distance, hex_to_rgb, mix
from ..parsing.occupancy import _layers, anchor_point, composite
from .base import AuditContext, check, finding

C = CheckCategory.template


@check("template.font_family", C, "Шрифт не из шаблона или гарнитур больше двух", fix="replace_font")
def font_family(ctx: AuditContext) -> list[Finding]:
    allowed = set(ctx.tm.design.fonts[: ctx.cfg.max_font_families])
    out, used = [], set()
    for sv in ctx.slides:
        for s in sv.text_shapes():
            fams = {r.family for r in s.runs}
            used |= fams
            bad = fams - allowed
            if bad:
                out.append(finding("template.font_family", sv.index, f"Гарнитура {sorted(bad)} не из шаблона {sorted(allowed)}",
                                   shape=s, severity=Severity.error))
    if len(used) > ctx.cfg.max_font_families:
        out.append(finding("template.font_family", 0, f"В колоде {len(used)} гарнитуры: {sorted(used)}", fix=None))
    return out


@check("template.font_size", C, "Кегль не из типографической шкалы шаблона", fix="snap_font_size")
def font_size(ctx: AuditContext) -> list[Finding]:
    d = ctx.tm.design
    scale = set(d.font_sizes) | {r.size_pt for r in d.type_scale}
    out = []
    for sv in ctx.slides:
        for s in sv.text_shapes():
            bad = sorted({r.size for r in s.runs if not any(abs(r.size - x) <= 0.26 for x in scale)})
            if bad:
                out.append(finding("template.font_size", sv.index, f"Кегль {bad} pt вне шкалы шаблона", shape=s, sizes=bad))
    return out


@check("template.color", C, "Цвет не из палитры шаблона", fix="off_palette_color")
def color(ctx: AuditContext) -> list[Finding]:
    pal = [c.hex for c in ctx.tm.design.palette]
    out = []

    def off(c: str | None) -> bool:
        return bool(c) and min(distance(c, p) for p in pal) > 28

    for sv in ctx.slides:
        for s in sv.shapes:
            cols = [r.color for r in s.runs] + ([s.fill] if s.fill_kind == "solid" else [])
            bad = sorted({c for c in cols if off(c)})
            if bad and (sv.is_new(s) or sv.is_changed(s) or s.kind == "sp" and s.text.strip() and s.id not in sv.ts.shape_boxes):
                out.append(finding("template.color", sv.index, f"Цвета {['#' + c for c in bad]} не из палитры", shape=s, colors=bad))
        ps = ctx.plan.slides[sv.index] if sv.index < len(ctx.plan.slides) else None
        if ps and ps.visual and ps.visual.chart:
            bad = sorted({x.color for x in ps.visual.chart.series if off(x.color)} | {c for c in ps.visual.chart.point_colors if off(c)})
            if bad:
                out.append(finding("template.color", sv.index, f"Цвета графика {['#' + c for c in bad]} не из палитры",
                                   box=ps.visual.box, colors=bad))
    return out


@check("template.layout", C, "Слайд собран не на макете из шаблона")
def layout(ctx: AuditContext) -> list[Finding]:
    names = set()
    try:
        from pptx import Presentation
        tpl = Presentation(ctx.plan.template_file)
        names = {lay.name for m in tpl.slide_masters for lay in m.slide_layouts}
    except Exception:  # noqa: BLE001
        return []
    return [finding("template.layout", sv.index, f"Макет «{sv.rs.layout_name}» отсутствует в шаблоне", severity=Severity.error)
            for sv in ctx.slides if sv.rs.layout_name not in names]


@check("template.logo_position", C, "Логотип или колонтитул сдвинуты с положенного места")
def logo_position(ctx: AuditContext) -> list[Finding]:
    out = []
    for sv in ctx.slides:
        logos = {p.shape_id: p for p in sv.ts.pictures if p.role == PictureRole.logo}
        foot = {t.shape_id for t in sv.ts.texts if t.role.value == "footer"}
        for s in sv.shapes:
            if (s.id in logos or s.id in foot or s.ph_type in ("ftr", "sldNum", "dt")) and sv.is_changed(s, eps=0.004) \
                    and s.id in sv.ts.shape_boxes:
                out.append(finding("template.logo_position", sv.index, "Логотип/колонтитул сдвинут относительно образца",
                                   shape=s, severity=Severity.error))
    return out


@check("template.contrast", C, "Текст плохо читается на фоне: APCA |Lc| < 45 (крупный) / 55 (обычный)", fix="low_contrast")
def contrast(ctx: AuditContext) -> list[Finding]:
    """APCA вместо WCAG 2: WCAG считает чёрный на фирменном синем #0077FF нормой (5.1:1), а белый — ошибкой (4.1:1),
    и автоисправление красило заголовки обложки в чёрный. Фон — итоговый цвет всех слоёв в точке начала текста
    (фон слайда, декор макета, картинки с прозрачностью, собственная плашка)."""
    out = []
    for sv in ctx.slides:
        img = _open(sv.png)
        for s in sv.text_shapes():
            bg = _backdrop(ctx, sv, s) or _background(sv, s, img)
            worst = None
            for r in s.runs:
                if not r.color or not bg or not r.text.strip():
                    continue
                lc = abs(apca_lc(r.color, bg))
                floor = 45.0 if r.size >= 24 or (r.bold and r.size >= 18) else 55.0
                if lc < floor and (worst is None or lc - floor < worst[0] - worst[2]):
                    worst = (lc, r.color, floor)
            if worst is None:
                continue
            tpl = sv.ts.slot(s.id)
            accent = not _neutral(worst[1])
            as_template = tpl is not None and tpl.color == worst[1] and not sv.is_new(s) and accent
            out.append(finding("template.contrast", sv.index,
                               f"APCA Lc {worst[0]:.0f} < {worst[2]:.0f}: #{worst[1]} на #{bg}"
                               + (" — акцентный цвет шаблона" if as_template else ""),
                               shape=s, severity=Severity.info if as_template else Severity.error,
                               fix="" if as_template else None,       # акцент дизайнера не перекрашиваем автоматически
                               lc=round(worst[0], 1), wcag=round(contrast_ratio(worst[1], bg), 2),
                               fg=worst[1], bg=bg, min_lc=worst[2], as_template=as_template))
    return out


def _neutral(h: str) -> bool:
    r, g, b = hex_to_rgb(h)
    return max(r, g, b) - min(r, g, b) < 48


def _backdrop(ctx: AuditContext, sv, s) -> str | None:
    if ctx.reader is None:
        return None
    below, _ = _layers(s, sv.shapes, sv.decor)
    ax, ay = anchor_point(s, s.text_box)
    try:
        return composite(ax, ay, sv.rs.background.color or "FFFFFF", below, ctx.reader)
    except Exception:  # noqa: BLE001 — битая картинка не должна ронять аудит
        return None


def _open(png):
    if png is None:
        return None
    try:
        from PIL import Image
        return Image.open(png).convert("RGB")
    except Exception:  # noqa: BLE001
        return None


def _background(sv, s, img) -> str | None:
    """Фон под текстом: по картинке рендера (кромка рамки — там нет букв), иначе заливка ближайшей подложки / фон слайда."""
    if img is not None:
        W, H = img.size
        b = s.text_box
        x0, y0 = max(0, int(b.x * W) - 3), max(0, int(b.y * H) - 3)
        x1, y1 = min(W - 1, int(b.x2 * W) + 3), min(H - 1, int(b.y2 * H) + 3)
        if x1 > x0 + 4 and y1 > y0 + 4:
            pts = [img.getpixel((x, y0)) for x in range(x0, x1, 3)] + [img.getpixel((x, y1)) for x in range(x0, x1, 3)]
            pts += [img.getpixel((x0, y)) for y in range(y0, y1, 3)] + [img.getpixel((x1, y)) for y in range(y0, y1, 3)]
            pts.sort(key=lambda p: sum(p))
            r, g, bb = pts[len(pts) // 2]
            return f"{r:02X}{g:02X}{bb:02X}"
    layers = [(0, o) for o in sv.decor] + [(1, o) for o in sv.shapes if o is not s and o.z < s.z]   # макет всегда под слайдом
    under = sorted(((lvl, o) for lvl, o in layers if o.fill and o.fill_kind in ("solid", "grad", "img") and o.box.area < 0.95
                    and o.box.contains_point(s.box.cx, s.box.cy)), key=lambda lo: (lo[0], lo[1].z))
    color = sv.rs.background.color or "FFFFFF"
    for _, o in under + ([(2, s)] if s.fill and s.fill_kind in ("solid", "grad", "img") else []):
        color = mix(color, o.fill, o.fill_alpha) if o.fill_alpha < 0.999 else o.fill   # полупрозрачная плашка — смесь с фоном
    return color
