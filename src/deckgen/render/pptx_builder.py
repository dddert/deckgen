"""DeckPlan -> .pptx: клон слайда-образца + правки. Все объекты нативные (текст, фигуры, таблицы, графики).

Почему клон, а не «пустой layout + нарисовать фигуры» (как в v1):
- в датасете дизайн живёт в слайдах-образцах (градиенты, скругления, иконки, декор, группы) —
  при перерисовке он терялся, при клонировании сохраняется на 100%;
- стиль текста (кегль/цвет/межстрочный/маркеры) берётся из рана-образца, а не пересобирается;
- слайд остаётся на своём макете шаблона и с его плейсхолдерами (проверка «слайд на макете шаблона»).
"""
from __future__ import annotations

import copy
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu, Pt

from ..models import Box, DeckPlan, ImageFill, PlannedSlide, TemplateModel, TextFill, TextSlot, VisualFill
from ..ooxml.reader import NS, q

_SHAPE_TAGS = {q("p:sp"), q("p:pic"), q("p:grpSp"), q("p:graphicFrame"), q("p:cxnSp")}
_R_ATTRS = {f"{{{NS['r']}}}{a}" for a in ("id", "embed", "link", "pict", "dm", "lo", "qs", "cs")}
_CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS, "area": XL_CHART_TYPE.AREA,
    "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT,
}


def build_pptx(plan: DeckPlan, tm: TemplateModel, out: Path, prune_layouts: bool = True) -> Path:
    prs = Presentation(plan.template_file)
    originals = list(prs.slides)
    for ps in plan.slides:
        src = originals[ps.template_slide]
        dst = clone_slide(prs, src)
        apply_plan(dst, ps, tm.slides[ps.template_slide], prs)
    for s in originals:
        _drop_slide(prs, s)
    if prune_layouts:
        _prune_layouts(prs)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


# ============================================================================ клонирование

def clone_slide(prs, src):
    dst = prs.slides.add_slide(src.slide_layout)
    tree = dst.shapes._spTree
    for el in list(tree):
        if el.tag in _SHAPE_TAGS:
            tree.remove(el)
    rid_map: dict[str, str] = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        if rel.is_external:
            rid_map[rid] = dst.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            rid_map[rid] = dst.part.relate_to(rel.target_part, rel.reltype)
    s_csld, d_csld = src._element.find(q("p:cSld")), dst._element.find(q("p:cSld"))
    bg = s_csld.find(q("p:bg"))
    if bg is not None:
        old = d_csld.find(q("p:bg"))
        if old is not None:
            d_csld.remove(old)
        d_csld.insert(0, copy.deepcopy(bg))
    if s_csld.get("name"):
        d_csld.set("name", s_csld.get("name"))
    for attr in ("showMasterSp", "showMasterPhAnim"):
        if src._element.get(attr) is not None:
            dst._element.set(attr, src._element.get(attr))
    ov = src._element.find(q("p:clrMapOvr"))
    if ov is not None:
        old = dst._element.find(q("p:clrMapOvr"))
        if old is not None:
            dst._element.replace(old, copy.deepcopy(ov))
    ext = tree.find(q("p:extLst"))
    for el in src.shapes._spTree:
        if el.tag in _SHAPE_TAGS or etree.QName(el).localname == "AlternateContent":
            c = copy.deepcopy(el)
            if ext is not None:
                ext.addprevious(c)
            else:
                tree.append(c)
    for root in (tree, d_csld.find(q("p:bg"))):
        if root is None:
            continue
        for el in root.iter():
            for a in list(el.attrib):
                if a in _R_ATTRS and el.get(a) in rid_map:
                    el.set(a, rid_map[el.get(a)])
    return dst


def _drop_slide(prs, slide) -> None:
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        if prs.part.related_part(sid.rId) is slide.part:
            prs.part.drop_rel(sid.rId)
            lst.remove(sid)
            return


def _prune_layouts(prs) -> None:
    """Неиспользованные макеты и мастера — в датасете это 10–17 МБ картинок, которые ни к чему."""
    used = {id(s.slide_layout.part) for s in prs.slides}
    for master in list(prs.slide_masters):
        keep = [lay for lay in master.slide_layouts if id(lay.part) in used]
        if not keep and len(prs.slide_masters) > 1:
            lst = prs.slide_masters._sldMasterIdLst
            for mid in list(lst):
                if prs.part.related_part(mid.rId) is master.part:
                    prs.part.drop_rel(mid.rId)
                    lst.remove(mid)
            continue
        for lay in list(master.slide_layouts):
            if id(lay.part) not in used and len(master.slide_layouts) > 1:
                master.slide_layouts.remove(lay)


# ============================================================================ правки

def _find(tree, shape_id: str):
    for el in tree.iter(*_SHAPE_TAGS):
        nv = next((c for c in el if etree.QName(c).localname.startswith("nv")), None)
        cnv = nv.find(q("p:cNvPr")) if nv is not None else None
        if cnv is not None and cnv.get("id") == shape_id:
            return el
    return None


def apply_plan(slide, ps: PlannedSlide, tslide, prs) -> None:
    tree = slide.shapes._spTree
    W, H = int(prs.slide_width), int(prs.slide_height)
    slots = {t.shape_id: t for t in tslide.texts}
    for sid, box in ps.moves.items():
        el = _find(tree, sid)
        if el is not None and sid in tslide.shape_boxes:
            _move(el, tslide.shape_boxes[sid], box, W, H)
    for tf in ps.texts:
        el = _find(tree, tf.shape_id)
        if el is not None:
            set_text(el, tf, slots.get(tf.shape_id))
            if tf.nowrap:
                bp = el.find(q("p:txBody")).find(q("a:bodyPr")) if el.find(q("p:txBody")) is not None else None
                if bp is not None:
                    bp.set("wrap", "none")
            if tf.fit_box is not None and tf.shape_id in tslide.shape_boxes:
                _set_box(el, tslide.shape_boxes[tf.shape_id], tf.fit_box, W, H)
            if tf.text_area is not None and tf.shape_id in tslide.shape_boxes:
                _set_insets(el, tslide.shape_boxes[tf.shape_id], tf.text_area, W, H)
    for im in ps.images:
        _set_image(slide, tree, im, tslide, W, H)
    if ps.visual is not None:
        _visual(slide, tree, ps.visual, W, H)
    for op in ps.style_ops:
        el = _find(tree, op.shape_id)
        if el is not None:
            _style(el, op)
    for sid in ps.delete_ids:
        el = _find(tree, sid)
        if el is not None:
            _remove(el)
    if ps.notes:
        slide.notes_slide.notes_text_frame.text = ps.notes


def _style(el, op) -> None:
    runs = [r.find(q("a:rPr")) for r in el.iter(q("a:r"))]
    runs = [r for r in runs if r is not None]
    if op.op == "size":
        for rpr in runs:
            rpr.set("sz", str(int(round(float(op.value) * 100))))
    elif op.op == "font":
        for rpr in runs:
            for tag in ("a:latin", "a:ea", "a:cs"):
                x = rpr.find(q(tag))
                if x is None:
                    x = etree.SubElement(rpr, q(tag))
                x.set("typeface", str(op.value))
    elif op.op == "recolor":
        if op.match and op.match.startswith("run:"):  # одна часть составного текста (синий заголовок карточки)
            k = int(op.match[4:])
            allr = [r for r in el.iter(q("a:r")) if (_run_text(r) or "").strip()]
            if k < len(allr):
                rpr = allr[k].find(q("a:rPr"))
                if rpr is None:
                    rpr = etree.Element(q("a:rPr"))
                    allr[k].insert(0, rpr)
                for old in list(rpr):
                    if etree.QName(old).localname in ("solidFill", "gradFill", "noFill"):
                        rpr.remove(old)
                sf = etree.Element(q("a:solidFill"))
                etree.SubElement(sf, q("a:srgbClr")).set("val", str(op.value).upper())
                ln = rpr.find(q("a:ln"))
                (ln.addnext(sf) if ln is not None else rpr.insert(0, sf))
        elif op.match:                                # конкретный цвет — во всей фигуре (текст и заливка)
            for c in el.iter(q("a:srgbClr")):
                if (c.get("val") or "").upper() == op.match.upper():
                    c.set("val", str(op.value).upper())
        else:                                         # весь текст фигуры — одним цветом
            for rpr in runs:
                for old in list(rpr):
                    if etree.QName(old).localname in ("solidFill", "gradFill", "noFill"):
                        rpr.remove(old)
                sf = etree.Element(q("a:solidFill"))
                etree.SubElement(sf, q("a:srgbClr")).set("val", str(op.value).upper())
                ln = rpr.find(q("a:ln"))
                (ln.addnext(sf) if ln is not None else rpr.insert(0, sf))


def _remove(el) -> None:
    parent = el.getparent()
    parent.remove(el)
    if parent.tag == q("p:grpSp") and not any(c.tag in _SHAPE_TAGS for c in parent):
        _remove(parent)


# ---------------------------------------------------------------------------- текст

def _segments(p) -> list[list]:
    """Строки абзаца: списки ранов, разделённые a:br."""
    segs, cur = [], []
    for c in p:
        tag = etree.QName(c).localname
        if tag == "br":
            segs.append(cur)
            cur = []
        elif tag in ("r", "fld"):
            cur.append(c)
    segs.append(cur)
    return segs


def _blank_run(p) -> etree._Element:
    r = etree.SubElement(p, q("a:r"))
    src = p.find(q("a:endParaRPr"))
    rpr = copy.deepcopy(src) if src is not None else etree.Element(q("a:rPr"))
    rpr.tag = q("a:rPr")
    r.append(rpr)
    etree.SubElement(r, q("a:t"))
    p.remove(r)
    return r


def _prep_run(r, text: str, size_pt: float | None):
    r = copy.deepcopy(r)
    if etree.QName(r).localname == "fld":
        nr = etree.Element(q("a:r"))
        for c in r:
            if etree.QName(c).localname in ("rPr", "t"):
                nr.append(c)
        r = nr
    rpr = r.find(q("a:rPr"))
    if rpr is None:
        rpr = etree.Element(q("a:rPr"))
        r.insert(0, rpr)
    for a in ("err", "dirty", "smtClean", "noProof"):
        rpr.attrib.pop(a, None)
    rpr.set("lang", "ru-RU")
    if size_pt:
        rpr.set("sz", str(int(round(size_pt * 100))))
    t = r.find(q("a:t"))
    if t is None:
        t = etree.SubElement(r, q("a:t"))
    t.text = text
    return r


def set_text(sp, tf: TextFill, slot: TextSlot | None) -> None:
    txb = sp.find(q("p:txBody"))
    if txb is None:
        return
    paras = txb.findall(q("a:p"))
    if not paras:
        paras = [etree.SubElement(txb, q("a:p"))]
    if slot is not None and slot.parts and len(slot.parts) > 1:
        _set_composite(txb, paras, tf)
        return
    filled = [p for p in paras if any(_run_text(c).strip() for c in p if etree.QName(c).localname in ("r", "fld"))]
    templates = filled or [p for p in paras if p.find(q("a:r")) is not None] or paras[:1]
    for p in paras:
        txb.remove(p)
    for i, text in enumerate(tf.paragraphs):
        tp = templates[min(i, len(templates) - 1)]
        # стиль берём у ранов с текстом: экспорт из Google Slides оставляет пустые раны с чёрным цветом по умолчанию
        runs = [c for c in tp if etree.QName(c).localname in ("r", "fld") and _run_text(c).strip()] \
            or [c for c in tp if etree.QName(c).localname in ("r", "fld")]
        br_tpl = tp.find(q("a:br"))
        p = copy.deepcopy(tp)
        for c in list(p):
            if etree.QName(c).localname in ("r", "fld", "br"):
                p.remove(c)
        end = p.find(q("a:endParaRPr"))
        lines = text.split("\n")
        for j, line in enumerate(lines):
            rt = runs[min(j, len(runs) - 1)] if runs else _blank_run(tp)
            if j > 0:
                br = copy.deepcopy(br_tpl) if br_tpl is not None else etree.Element(q("a:br"))
                if br.find(q("a:rPr")) is None and rt.find(q("a:rPr")) is not None:
                    br.append(copy.deepcopy(rt.find(q("a:rPr"))))
                _insert(p, br, end)
            _insert(p, _prep_run(rt, line, tf.size_pt), end)
        txb.append(p)


def _run_text(r) -> str:
    t = r.find(q("a:t"))
    return (t.text or "") if t is not None else ""


def _set_composite(txb, paras, tf: TextFill) -> None:
    """«ХХ% + подпись»: часть i пишется в i-й сегмент образца с его собственным стилем."""
    segs = [(p, seg) for p in paras for seg in _segments(p) if "".join(
        (c.find(q("a:t")).text or "") for c in seg if c.find(q("a:t")) is not None).strip()]
    for i, (p, seg) in enumerate(segs):
        if i < len(tf.paragraphs) and tf.paragraphs[i].strip():
            first = max(seg, key=lambda c: len(_run_text(c).strip()))
            for c in seg:
                if c is not first:
                    p.remove(c)
            seg = [first]
            new = _prep_run(first, tf.paragraphs[i], None)
            p.replace(first, new)
        else:
            for c in seg:
                prev = c.getprevious()
                if prev is not None and etree.QName(prev).localname == "br":
                    p.remove(prev)
                p.remove(c)
    for p in paras:
        if not any(etree.QName(c).localname in ("r", "fld") for c in p) and len(txb.findall(q("a:p"))) > 1:
            txb.remove(p)


def _insert(p, el, end) -> None:
    if end is not None:
        end.addprevious(el)
    else:
        p.append(el)


# ---------------------------------------------------------------------------- геометрия

def _group_scale(el) -> tuple[float, float]:
    sx = sy = 1.0
    parent = el.getparent()
    while parent is not None and parent.tag == q("p:grpSp"):
        xf = parent.find(q("p:grpSpPr")).find(q("a:xfrm"))
        if xf is not None and xf.find(q("a:chExt")) is not None:
            ext, chx = xf.find(q("a:ext")), xf.find(q("a:chExt"))
            if int(chx.get("cx", 0)):
                sx *= int(ext.get("cx", 0)) / int(chx.get("cx"))
            if int(chx.get("cy", 0)):
                sy *= int(ext.get("cy", 0)) / int(chx.get("cy"))
        parent = parent.getparent()
    return sx or 1.0, sy or 1.0


def _xfrm(el):
    if el.tag == q("p:grpSp"):
        return el.find(q("p:grpSpPr")).find(q("a:xfrm"))
    if el.tag == q("p:graphicFrame"):
        return el.find(q("p:xfrm"))
    sppr = el.find(q("p:spPr"))
    return sppr.find(q("a:xfrm")) if sppr is not None else None


def _move(el, old: Box, new: Box, W: int, H: int) -> None:
    xf = _xfrm(el)
    if xf is None or xf.find(q("a:off")) is None:
        return
    sx, sy = _group_scale(el)
    off = xf.find(q("a:off"))
    off.set("x", str(int(int(off.get("x")) + (new.x - old.x) * W / sx)))
    off.set("y", str(int(int(off.get("y")) + (new.y - old.y) * H / sy)))


def _set_box(el, old: Box, new: Box, W: int, H: int) -> None:
    """Новая рамка фигуры (расширенный заголовок, выросший по тексту блок) — в координатах её группы."""
    xf = _xfrm(el)
    if xf is None or xf.find(q("a:ext")) is None or xf.find(q("a:off")) is None or old.w <= 0 or old.h <= 0:
        return
    sx, sy = _group_scale(el)
    off, ext = xf.find(q("a:off")), xf.find(q("a:ext"))
    off.set("x", str(int(int(off.get("x")) + (new.x - old.x) * W / sx)))
    off.set("y", str(int(int(off.get("y")) + (new.y - old.y) * H / sy)))
    ext.set("cx", str(max(1, int(int(ext.get("cx")) * new.w / old.w))))
    ext.set("cy", str(max(1, int(int(ext.get("cy")) * new.h / old.h))))


def _set_insets(el, shape: Box, area: Box, W: int, H: int) -> None:
    """Текст — в свободной части плашки: внутренние поля вместо сдвига самой плашки."""
    tx = el.find(q("p:txBody"))
    bp = tx.find(q("a:bodyPr")) if tx is not None else None
    if bp is None:
        return
    for attr, v in (("lIns", (area.x - shape.x) * W), ("tIns", (area.y - shape.y) * H),
                    ("rIns", (shape.x2 - area.x2) * W), ("bIns", (shape.y2 - area.y2) * H)):
        bp.set(attr, str(max(0, int(v))))


def _emu(b: Box, W: int, H: int) -> tuple[Emu, Emu, Emu, Emu]:
    return Emu(int(b.x * W)), Emu(int(b.y * H)), Emu(int(b.w * W)), Emu(int(b.h * H))


# ---------------------------------------------------------------------------- картинки

def _set_image(slide, tree, im: ImageFill, tslide, W: int, H: int) -> None:
    el = _find(tree, im.shape_id)
    if el is None:
        return
    box = tslide.shape_boxes.get(im.shape_id)
    if el.tag == q("p:pic"):
        _, rid = slide.part.get_or_add_image_part(im.path)
        blip = el.find(".//" + q("a:blip"))
        blip.set(q("r:embed"), rid)
        for c in list(blip):
            blip.remove(c)                       # эффекты образца (duotone и т.п.) к новой картинке не применимы
        _crop_fill(el.find(q("p:blipFill")), im.path, box, W, H)
        return
    for ph in slide.placeholders:               # плейсхолдер-картинка: insert_picture сам обрежет под рамку
        if str(ph.shape_id) == im.shape_id and hasattr(ph, "insert_picture"):
            ph.insert_picture(im.path)
            return
    if box is None:
        return
    x, y, w, h = _emu(box, W, H)
    pic = slide.shapes.add_picture(im.path, x, y, w, h)
    _crop_fill(pic._element.find(q("p:blipFill")), im.path, box, W, H)
    el.addprevious(pic._element)                # на место заглушки по z-порядку
    _remove(el)


def _crop_fill(blip_fill, path: str, box: Box | None, W: int, H: int) -> None:
    """Картинка заполняет рамку без искажения пропорций (srcRect)."""
    if blip_fill is None or box is None:
        return
    from PIL import Image
    with Image.open(path) as im:
        src = im.width / im.height
    dst = (box.w * W) / max(1, box.h * H)
    rect = blip_fill.find(q("a:srcRect"))
    if rect is None:
        rect = etree.Element(q("a:srcRect"))
        blip = blip_fill.find(q("a:blip"))
        blip.addnext(rect)
    for k in ("l", "t", "r", "b"):
        rect.attrib.pop(k, None)
    if src > dst:
        k = int((1 - dst / src) / 2 * 100000)
        rect.set("l", str(k))
        rect.set("r", str(k))
    elif src < dst:
        k = int((1 - src / dst) / 2 * 100000)
        rect.set("t", str(k))
        rect.set("b", str(k))


# ---------------------------------------------------------------------------- графики и таблицы

def _visual(slide, tree, v: VisualFill, W: int, H: int) -> None:
    anchor = None
    for sid in v.remove_ids:
        el = _find(tree, sid)
        if el is not None:
            anchor = anchor if anchor is not None else el
    if v.kind == "table" and v.table_shape_id:
        el = _find(tree, v.table_shape_id)
        if el is not None and v.table is not None:
            _fill_template_table(el, v, W, H)
            return
    x, y, w, h = _emu(v.box, W, H)
    if v.kind == "chart" and v.chart is not None:
        frame = _add_chart(slide, v, x, y, w, h)
    elif v.table is not None:
        frame = _add_table(slide, v, x, y, w, h)
    else:
        return
    if anchor is not None:
        anchor.addprevious(frame._element)
    for sid in v.remove_ids:
        el = _find(tree, sid)
        if el is not None:
            _remove(el)


def _rgb(h: str | None) -> RGBColor | None:
    return RGBColor.from_string(h.upper()) if h and len(h) == 6 else None


def _add_chart(slide, v: VisualFill, x, y, w, h):
    spec = v.chart
    data = CategoryChartData(number_format=spec.number_format)
    data.categories = spec.categories
    for s in spec.series:
        data.add_series(s.name, s.values)
    frame = slide.shapes.add_chart(_CHART_TYPES[spec.chart_type], x, y, w, h, data)
    chart = frame.chart
    chart.has_title = False            # иначе PowerPoint подставит имя единственной серии как заголовок
    text = _rgb(v.text_color)
    chart.font.name = v.font_family or None
    chart.font.size = Pt(v.font_size or 12)
    if text is not None:
        chart.font.color.rgb = text
    pie = spec.chart_type in ("pie", "doughnut")
    chart.has_legend = spec.legend
    if spec.legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(v.font_size or 12)
        if text is not None:
            chart.legend.font.color.rgb = text
    plot = chart.plots[0]
    if pie:
        plot.vary_by_categories = True
        colors = spec.point_colors or [s.color for s in spec.series if s.color]
        for i, pt in enumerate(plot.series[0].points):
            c = _rgb(colors[i % len(colors)]) if colors else None
            if c is not None:
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = c
    else:
        plot.vary_by_categories = False
        for i, s in enumerate(spec.series):
            c = _rgb(s.color)
            ser = plot.series[i]
            if c is None:
                continue
            if spec.chart_type in ("line",):
                ser.format.line.color.rgb = c
                ser.format.line.width = Pt(2.5)
                ser.smooth = False
            else:
                ser.format.fill.solid()
                ser.format.fill.fore_color.rgb = c
        if spec.chart_type in ("column", "bar"):
            plot.gap_width = 60
            plot.overlap = -10 if len(spec.series) > 1 else 0
        grid = _rgb(v.grid_color)
        va, ca = chart.value_axis, chart.category_axis
        for ax in (va, ca):
            ax.tick_labels.font.size = Pt(v.font_size or 12)
            if text is not None:
                ax.tick_labels.font.color.rgb = text
            ax.format.line.fill.background() if ax is va else None
        va.has_major_gridlines = True
        if grid is not None:
            va.major_gridlines.format.line.color.rgb = grid
            va.major_gridlines.format.line.width = Pt(0.75)
            ca.format.line.color.rgb = grid
        if spec.x_title:
            ca.has_title = True
            ca.axis_title.text_frame.text = spec.x_title
            _style_title(ca.axis_title, v, text)
        if spec.y_title:
            va.has_title = True
            va.axis_title.text_frame.text = spec.y_title
            _style_title(va.axis_title, v, text)
    if spec.data_labels:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.font.size = Pt(max(9, (v.font_size or 12) - 1))
        if text is not None and spec.chart_type not in ("pie", "doughnut"):
            dl.font.color.rgb = text
        elif pie:
            dl.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        dl.number_format = spec.number_format
        dl.number_format_is_linked = spec.number_format == "General"
        if spec.chart_type in ("column", "bar"):
            dl.position = XL_LABEL_POSITION.OUTSIDE_END
        elif spec.chart_type == "line":
            dl.position = XL_LABEL_POSITION.ABOVE
    # прозрачный фон области графика — лежит на фоне слайда шаблона
    cs = chart._chartSpace
    _no_fill(cs, "c:spPr")
    pa = cs.find(".//" + q("c:plotArea"))
    if pa is not None:
        _no_fill(pa, "c:spPr")
    return frame


def _style_title(title, v: VisualFill, text) -> None:
    for p in title.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(max(9, (v.font_size or 12) - 1))
            r.font.bold = False
            if v.font_family:
                r.font.name = v.font_family
            if text is not None:
                r.font.color.rgb = text


def _no_fill(parent, tag: str) -> None:
    sppr = parent.find(q(tag))
    if sppr is None:
        sppr = etree.SubElement(parent, q(tag))
    for c in list(sppr):
        if etree.QName(c).localname in ("solidFill", "noFill", "gradFill", "ln"):
            sppr.remove(c)
    etree.SubElement(sppr, q("a:noFill"))
    ln = etree.SubElement(sppr, q("a:ln"))
    etree.SubElement(ln, q("a:noFill"))


def _add_table(slide, v: VisualFill, x, y, w, h):
    t = v.table
    rows, cols = len(t.rows) + 1, len(t.columns)
    row_h = Emu(int(min(h / rows, Pt((v.font_size or 12) * 2.6))))
    frame = slide.shapes.add_table(rows, cols, x, y, w, Emu(row_h * rows))
    tbl = frame.table
    text = _rgb(v.text_color)
    head = _rgb(v.header_fill)
    for r in range(rows):
        tbl.rows[r].height = row_h
        for c in range(cols):
            cell = tbl.cell(r, c)
            val = t.columns[c] if r == 0 else (t.rows[r - 1][c] if c < len(t.rows[r - 1]) else "")
            cell.text = str(val)
            cell.margin_left = cell.margin_right = Pt(6)
            cell.margin_top = cell.margin_bottom = Pt(3)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(v.font_size or 12)
                    run.font.bold = r == 0
                    if v.font_family:
                        run.font.name = v.font_family
                    if r == 0 and head is not None:
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    elif text is not None:
                        run.font.color.rgb = text
            if r == 0 and head is not None:
                cell.fill.solid()
                cell.fill.fore_color.rgb = head
            else:
                cell.fill.background()
    # у встроенного стиля таблицы python-pptx — синие полосы; отключаем, оставляем тонкие линии
    tblpr = frame._element.find(".//" + q("a:tblPr"))
    if tblpr is not None:
        tblpr.set("bandRow", "0")
        tblpr.set("firstRow", "1")
    return frame


def _fill_template_table(frame_el, v: VisualFill, W: int, H: int) -> None:
    """Нативная таблица образца: строки клонируются со стилем шапки/тела, колонки подгоняются под данные."""
    t = v.table
    tbl = frame_el.find(".//" + q("a:tbl"))
    grid = tbl.find(q("a:tblGrid"))
    trs = tbl.findall(q("a:tr"))
    if not trs:
        return
    head_tpl = trs[0]
    body_tpls = trs[1:3] or trs[:1]
    ncols = len(t.columns)
    gcols = grid.findall(q("a:gridCol"))
    total_w = sum(int(g.get("w", 0)) for g in gcols)
    for g in gcols:
        grid.remove(g)
    for _ in range(ncols):
        g = copy.deepcopy(gcols[0])
        g.set("w", str(int(total_w / ncols)))
        grid.append(g)
    for tr in trs:
        tbl.remove(tr)

    def make_row(tpl, values: list[str]):
        tr = copy.deepcopy(tpl)
        tcs = tr.findall(q("a:tc"))
        for tc in tcs:
            tr.remove(tc)
        for i, val in enumerate(values):
            tc = copy.deepcopy(tcs[min(i, len(tcs) - 1)])
            for a in ("gridSpan", "hMerge", "vMerge", "rowSpan"):
                tc.attrib.pop(a, None)
            txb = tc.find(q("a:txBody"))
            paras = txb.findall(q("a:p"))
            tp = next((p for p in paras if p.find(q("a:r")) is not None), paras[0])
            for p in paras:
                txb.remove(p)
            p = copy.deepcopy(tp)
            runs = [c for c in p if etree.QName(c).localname in ("r", "fld", "br")]
            rt = next((c for c in runs if etree.QName(c).localname == "r"), None)
            for c in runs:
                p.remove(c)
            end = p.find(q("a:endParaRPr"))
            _insert(p, _prep_run(rt if rt is not None else _blank_run(p), str(val), None), end)
            txb.append(p)
            tr.append(tc)
        return tr

    tbl.append(make_row(head_tpl, [str(c) for c in t.columns]))
    for i, row in enumerate(t.rows):
        vals = [str(row[c]) if c < len(row) else "" for c in range(ncols)]
        tbl.append(make_row(body_tpls[i % len(body_tpls)], vals))
    # высота рамки = сумма строк, но не больше места под таблицу
    heights = [int(tr.get("h", 0)) for tr in tbl.findall(q("a:tr"))]
    xf = frame_el.find(q("p:xfrm"))
    limit = int(v.box.h * H)
    if sum(heights) > limit and heights:
        k = limit / sum(heights)
        for tr in tbl.findall(q("a:tr")):
            tr.set("h", str(int(int(tr.get("h", 0)) * k)))
    if xf is not None:
        xf.find(q("a:ext")).set("cy", str(min(limit, sum(int(tr.get("h", 0)) for tr in tbl.findall(q("a:tr"))))))
