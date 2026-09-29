"""OOXML -> «разрешённые» фигуры слайда: абсолютная геометрия и вычисленные стили.

Один и тот же ридер читает и шаблон (parsing), и готовый файл (audit, html). Что он делает,
чего не делал v1:
- трансформации групп (chOff/chExt): в датасете 44 группы, где координаты детей не совпадают
  с координатами на слайде;
- текст группы = текст её детей, а не их повтор (v1 считал раны детей дважды);
- наследование плейсхолдеров slide -> layout -> master: позиция, кегль, гарнитура, цвет, маркеры
  списка, интерлиньяж, автоподбор (≈10% ранов в датасете без явного кегля);
- цвета через тему конкретного мастера и его clrMap; фон слайда/макета/мастера.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from ..models import Box
from .colors import DEFAULT_CLRMAP, ColorContext, Theme

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
}
A, P, R = NS["a"], NS["p"], NS["r"]
EMU_PER_IN = 914400
EMU_PER_PT = 12700
_SHAPE_TAGS = {"sp", "pic", "grpSp", "graphicFrame", "cxnSp"}
_DEFAULT_INS = (91440, 45720, 91440, 45720)   # l, t, r, b


def q(tag: str) -> str:
    pfx, name = tag.split(":")
    return f"{{{NS[pfx]}}}{name}"


@dataclass
class RRun:
    text: str
    size: float
    family: str
    bold: bool = False
    italic: bool = False
    color: str | None = None


@dataclass
class RPara:
    runs: list[RRun]
    level: int = 0
    bullet: bool = False
    align: str = "l"
    line_spacing: float = 1.0

    @property
    def text(self) -> str:
        return "".join(r.text for r in self.runs)


@dataclass
class RShape:
    id: str
    name: str
    kind: str                      # sp | pic | grp | frame | cxn
    box: Box
    rot: float = 0.0
    flip: tuple[bool, bool] = (False, False)   # flipH, flipV — зеркальная картинка (свечение VK WorkSpace)
    ph_type: str | None = None
    ph_idx: str | None = None
    geometry: str | None = None
    fill: str | None = None
    fill_kind: str | None = None   # solid | grad | img | none
    fill_alpha: float = 1.0        # непрозрачность заливки (плашка 20% поверх фона)
    grad: list[str] = field(default_factory=list)
    line: str | None = None
    paras: list[RPara] = field(default_factory=list)
    image_part: str | None = None
    image_px: tuple[int, int] | None = None
    crop: tuple[float, float, float, float] = (0, 0, 0, 0)   # l, t, r, b (доли)
    table: list[list[str]] | None = None
    table_font: RRun | None = None
    chart_part: str | None = None
    diagram: bool = False
    parent: str | None = None
    top: str = ""
    depth: int = 0
    z: int = 0
    autofit: str = "none"          # none | shape | norm
    insets: tuple[float, float, float, float] = (0, 0, 0, 0)   # доли слайда
    anchor: str = "t"
    wrap: bool = True              # bodyPr wrap="none" — текст не переносится, выходит за рамку по ширине
    txbox: bool = False
    default_run: RRun | None = None   # стиль, который получит текст пустого плейсхолдера (наследуется от макета)
    el: etree._Element | None = field(default=None, repr=False)

    @property
    def text(self) -> str:
        return "\n".join(p.text for p in self.paras).strip()

    @property
    def runs(self) -> list[RRun]:
        return [r for p in self.paras for r in p.runs if r.text.strip()]

    @property
    def main_run(self) -> RRun | None:
        runs = self.runs
        return max(runs, key=lambda r: len(r.text)) if runs else self.default_run

    @property
    def text_box(self) -> Box:
        l, t, r, b = self.insets
        return Box(x=self.box.x + l, y=self.box.y + t, w=max(1e-4, self.box.w - l - r), h=max(1e-4, self.box.h - t - b))


@dataclass
class Background:
    kind: str = "none"             # solid | grad | img | none
    color: str | None = None       # сплошной цвет или средний цвет градиента/картинки
    image_part: str | None = None
    grad: list[str] = field(default_factory=list)


@dataclass
class RSlide:
    index: int
    partname: str
    layout_part: str
    layout_name: str
    master_part: str
    shapes: list[RShape]
    background: Background
    notes: str = ""

    def by_id(self) -> dict[str, RShape]:
        return {s.id: s for s in self.shapes}


class _Xf:
    """Аффинное преобразование EMU ребёнка группы -> EMU слайда (без поворота)."""

    def __init__(self, sx=1.0, sy=1.0, tx=0.0, ty=0.0):
        self.sx, self.sy, self.tx, self.ty = sx, sy, tx, ty

    def then_group(self, xfrm: etree._Element | None) -> "_Xf":
        if xfrm is None:
            return self
        off, ext = xfrm.find(q("a:off")), xfrm.find(q("a:ext"))
        choff, chext = xfrm.find(q("a:chOff")), xfrm.find(q("a:chExt"))
        if off is None or ext is None:
            return self
        cx, cy = int(ext.get("cx", 0)), int(ext.get("cy", 0))
        ccx = int(chext.get("cx", 0)) if chext is not None else cx
        ccy = int(chext.get("cy", 0)) if chext is not None else cy
        gsx = cx / ccx if ccx else 1.0
        gsy = cy / ccy if ccy else 1.0
        gtx = int(off.get("x", 0)) - (int(choff.get("x", 0)) if choff is not None else 0) * gsx
        gty = int(off.get("y", 0)) - (int(choff.get("y", 0)) if choff is not None else 0) * gsy
        return _Xf(gsx * self.sx, gsy * self.sy, gtx * self.sx + self.tx, gty * self.sy + self.ty)

    def apply(self, x: int, y: int, w: int, h: int) -> tuple[float, float, float, float]:
        return x * self.sx + self.tx, y * self.sy + self.ty, w * self.sx, h * self.sy


class PresentationReader:
    def __init__(self, source: str | Path | object):
        self.prs = Presentation(str(source)) if isinstance(source, (str, Path)) else source
        self.W, self.H = int(self.prs.slide_width), int(self.prs.slide_height)
        self._themes: dict[str, Theme] = {}
        self._img_px: dict[str, tuple[int, int]] = {}
        self._img_avg: dict[str, str] = {}
        self._defaults = self.prs.part._element.find(q("p:defaultTextStyle"))

    # ------------------------------------------------------------------ public
    @property
    def slide_w_in(self) -> float:
        return self.W / EMU_PER_IN

    @property
    def slide_h_in(self) -> float:
        return self.H / EMU_PER_IN

    def slides(self) -> list[RSlide]:
        return [self.read_slide(i) for i in range(len(self.prs.slides))]

    def read_slide(self, index: int) -> RSlide:
        slide = self.prs.slides[index]
        layout = slide.slide_layout
        master = layout.slide_master
        ctx = self._ctx(slide, layout, master)
        env = _Env(self, ctx, slide.part, layout, master)
        shapes: list[RShape] = []
        self._walk(slide._element.find(q("p:cSld")).find(q("p:spTree")), _Xf(), env, shapes, None, None, 0)
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else ""
        return RSlide(index=index, partname=str(slide.part.partname).lstrip("/"),
                      layout_part=str(layout.part.partname).lstrip("/"), layout_name=layout.name or "",
                      master_part=str(master.part.partname).lstrip("/"), shapes=shapes,
                      background=self.background(slide, layout, master, ctx), notes=notes)

    def read_decor(self, index: int) -> list[RShape]:
        """Не-плейсхолдерные фигуры мастера и макета под слайдом (логотипы, декор) — для HTML и контраста."""
        slide = self.prs.slides[index]
        layout = slide.slide_layout
        master = layout.slide_master
        if slide._element.get("showMasterSp") == "0":
            return []                               # «скрыть фоновую графику» на слайде
        ctx = self._ctx(slide, layout, master)
        owners = [layout] if layout._element.get("showMasterSp") == "0" else [master, layout]
        out: list[RShape] = []
        for owner in owners:
            env = _Env(self, ctx, owner.part, layout, master)
            tmp: list[RShape] = []
            self._walk(owner._element.find(q("p:cSld")).find(q("p:spTree")), _Xf(), env, tmp, None, None, 0)
            out += [s for s in tmp if not s.ph_type]
        return out

    def image(self, partname: str, max_px: int = 360):
        """Картинка части (RGBA, уменьшенная) — для карты занятости под текстом и оценки фона. None, если не растр."""
        key = (partname, max_px)
        if not hasattr(self, "_imgs"):
            self._imgs: dict = {}
            self._parts = {str(p.partname).lstrip("/"): p for p in self.prs.part.package.iter_parts()}
        if key not in self._imgs:
            part = self._parts.get(partname)
            im = None
            if part is not None:
                try:
                    from PIL import Image, ImageFilter
                    with Image.open(io.BytesIO(part.blob)) as src:
                        im = src.convert("RGBA").convert("RGBa")        # премультипликация: края без тёмного ореола
                        if max(im.size) > max_px:
                            k = max_px / max(im.size)
                            im = im.resize((max(1, int(im.width * k)), max(1, int(im.height * k))),
                                           Image.Resampling.BOX)
                        # зерно/дизеринг градиентных фонов — не край объекта: сглаживаем на размер клетки сетки
                        im = im.filter(ImageFilter.GaussianBlur(1.2)).convert("RGBA")
                except Exception:  # noqa: BLE001 — EMF/SVG
                    im = None
            self._imgs[key] = im
        return self._imgs[key]

    def image_px(self, part) -> tuple[int, int] | None:
        name = str(part.partname)
        if name not in self._img_px:
            try:
                from PIL import Image
                with Image.open(io.BytesIO(part.blob)) as im:
                    self._img_px[name] = im.size
            except Exception:  # noqa: BLE001 — EMF/WMF/SVG PIL не читает
                self._img_px[name] = (0, 0)
        v = self._img_px[name]
        return v if v != (0, 0) else None

    def image_avg(self, part) -> str | None:
        name = str(part.partname)
        if name not in self._img_avg:
            try:
                from PIL import Image
                with Image.open(io.BytesIO(part.blob)) as im:
                    r, g, b = im.convert("RGB").resize((1, 1), Image.Resampling.BOX).getpixel((0, 0))
                    self._img_avg[name] = f"{r:02X}{g:02X}{b:02X}"
            except Exception:  # noqa: BLE001
                self._img_avg[name] = ""
        return self._img_avg[name] or None

    def theme(self, master) -> Theme:
        key = str(master.part.partname)
        if key not in self._themes:
            try:
                blob = master.part.part_related_by(RT.THEME).blob
                self._themes[key] = Theme.from_xml(blob)
            except (KeyError, etree.XMLSyntaxError):
                self._themes[key] = Theme()
        return self._themes[key]

    def background(self, slide, layout, master, ctx: ColorContext) -> Background:
        for owner in (slide, layout, master):
            bg = owner._element.find(q("p:cSld")).find(q("p:bg"))
            if bg is None:
                continue
            pr = bg.find(q("p:bgPr"))
            if pr is not None:
                if pr.find(q("a:solidFill")) is not None:
                    return Background("solid", ctx.resolve(pr.find(q("a:solidFill"))))
                if pr.find(q("a:gradFill")) is not None:
                    stops = [ctx.resolve(gs) for gs in pr.iter(q("a:gs"))]
                    stops = [s for s in stops if s]
                    return Background("grad", _avg(stops), grad=stops)
                blip = pr.find(q("a:blipFill"))
                if blip is not None:
                    part = _related(owner.part, blip.find(q("a:blip")))
                    if part is not None:
                        return Background("img", self.image_avg(part), image_part=str(part.partname).lstrip("/"))
            ref = bg.find(q("p:bgRef"))
            if ref is not None:
                return Background("solid", ctx.resolve(ref))
        return Background("solid", ctx.scheme("bg1") or "FFFFFF")

    # ---------------------------------------------------------------- internals
    def _ctx(self, slide, layout, master) -> ColorContext:
        cmap = dict(DEFAULT_CLRMAP)
        mm = master._element.find(q("p:clrMap"))
        if mm is not None:
            cmap.update(dict(mm.attrib))
        for owner in (layout, slide):
            ov = owner._element.find(q("p:clrMapOvr"))
            if ov is not None:
                o = ov.find(q("a:overrideClrMapping"))
                if o is not None:
                    cmap.update(dict(o.attrib))
        return ColorContext(self.theme(master), cmap)

    def _walk(self, tree, xf: _Xf, env: "_Env", out: list[RShape], parent: str | None, top: str | None, depth: int):
        for el in tree:
            tag = etree.QName(el).localname
            if tag == "AlternateContent":
                branch = el.find(q("mc:Choice"))
                if branch is None or not any(etree.QName(c).localname in _SHAPE_TAGS for c in branch):
                    branch = el.find(q("mc:Fallback"))
                if branch is not None:
                    self._walk(branch, xf, env, out, parent, top, depth)
                continue
            if tag not in _SHAPE_TAGS:
                continue
            s = self._shape(el, tag, xf, env)
            if s is None:
                continue
            s.parent, s.depth, s.z = parent, depth, len(out)
            s.top = top or s.id
            out.append(s)
            if tag == "grpSp":
                gx = xf.then_group(el.find(q("p:grpSpPr")).find(q("a:xfrm")))
                self._walk(el, gx, env, out, s.id, top or s.id, depth + 1)

    def _shape(self, el, tag: str, xf: _Xf, env: "_Env") -> RShape | None:
        nv = next((c for c in el if etree.QName(c).localname.startswith("nv")), None)
        cnv = nv.find(q("p:cNvPr")) if nv is not None else None
        if cnv is None or cnv.get("hidden") in ("1", "true"):
            return None
        ph = nv.find(q("p:nvPr")).find(q("p:ph")) if nv.find(q("p:nvPr")) is not None else None
        ph_type = (ph.get("type") or "body") if ph is not None else None
        ph_idx = ph.get("idx") if ph is not None else None
        inh = env.inherited(ph_type, ph_idx) if ph is not None else []

        if tag == "grpSp":
            xfrm = el.find(q("p:grpSpPr")).find(q("a:xfrm"))
        elif tag == "graphicFrame":
            xfrm = el.find(q("p:xfrm"))
        else:
            sppr = el.find(q("p:spPr"))
            xfrm = sppr.find(q("a:xfrm")) if sppr is not None else None
        if (xfrm is None or xfrm.find(q("a:off")) is None) and inh:
            for base in inh:
                bx = base.find(q("p:spPr"))
                bx = bx.find(q("a:xfrm")) if bx is not None else None
                if bx is not None and bx.find(q("a:off")) is not None:
                    xfrm = bx
                    break
        if xfrm is None or xfrm.find(q("a:off")) is None or xfrm.find(q("a:ext")) is None:
            return None
        off, ext = xfrm.find(q("a:off")), xfrm.find(q("a:ext"))
        X, Y, Wd, Ht = xf.apply(int(off.get("x", 0)), int(off.get("y", 0)), int(ext.get("cx", 0)), int(ext.get("cy", 0)))
        box = Box(x=X / self.W, y=Y / self.H, w=Wd / self.W, h=Ht / self.H)
        kind = {"sp": "sp", "pic": "pic", "grpSp": "grp", "graphicFrame": "frame", "cxnSp": "cxn"}[tag]
        s = RShape(id=cnv.get("id", ""), name=cnv.get("name", ""), kind=kind, box=box,
                   rot=int(xfrm.get("rot", 0)) / 60000, ph_type=ph_type, ph_idx=ph_idx, el=el,
                   flip=(xfrm.get("flipH") in ("1", "true"), xfrm.get("flipV") in ("1", "true")))

        sppr = el.find(q("p:spPr"))
        if sppr is not None:
            geom = sppr.find(q("a:prstGeom"))
            s.geometry = geom.get("prst") if geom is not None else ("custom" if sppr.find(q("a:custGeom")) is not None else None)
            self._fill(s, sppr, el, env)
            ln = sppr.find(q("a:ln"))
            if ln is not None and ln.find(q("a:noFill")) is None:
                s.line = env.ctx.resolve(ln.find(q("a:solidFill")))
        if tag == "sp":
            txb = el.find(q("p:txBody"))
            nvsp = nv.find(q("p:cNvSpPr"))
            s.txbox = nvsp is not None and nvsp.get("txBox") == "1"
            if txb is not None:
                self._text(s, txb, ph_type, inh, env)
        elif tag == "pic":
            blip = el.find(q("p:blipFill")).find(q("a:blip")) if el.find(q("p:blipFill")) is not None else None
            part = _related(env.part, blip)
            if part is not None:
                s.image_part = str(part.partname).lstrip("/")
                s.image_px = self.image_px(part)
            src = el.find(q("p:blipFill")).find(q("a:srcRect")) if el.find(q("p:blipFill")) is not None else None
            if src is not None:
                s.crop = tuple(int(src.get(k, 0)) / 100000 for k in ("l", "t", "r", "b"))  # type: ignore[assignment]
        elif tag == "graphicFrame":
            gd = el.find(".//" + q("a:graphicData"))
            uri = gd.get("uri", "") if gd is not None else ""
            if uri.endswith("/table"):
                self._table(s, gd, env)
            elif uri.endswith("/chart"):
                ch = gd.find(q("c:chart"))
                part = _related(env.part, ch, attr=q("r:id"))
                s.chart_part = str(part.partname).lstrip("/") if part is not None else "chart"
            elif "diagram" in uri:
                s.diagram = True
        return s

    def _fill(self, s: RShape, sppr, el, env: "_Env") -> None:
        if sppr.find(q("a:noFill")) is not None:
            s.fill_kind = "none"
        elif sppr.find(q("a:solidFill")) is not None:
            s.fill_kind, s.fill = "solid", env.ctx.resolve(sppr.find(q("a:solidFill")))
            a = sppr.find(q("a:solidFill")).find(".//" + q("a:alpha"))
            if a is not None:
                s.fill_alpha = int(a.get("val", 100000)) / 100000
        elif sppr.find(q("a:gradFill")) is not None:
            stops = [env.ctx.resolve(gs) for gs in sppr.find(q("a:gradFill")).iter(q("a:gs"))]
            s.grad = [c for c in stops if c]
            s.fill_kind, s.fill = "grad", _avg(s.grad)
        elif sppr.find(q("a:blipFill")) is not None:
            s.fill_kind = "img"
            part = _related(env.part, sppr.find(q("a:blipFill")).find(q("a:blip")))
            if part is not None:
                s.image_part = str(part.partname).lstrip("/")
                s.fill = self.image_avg(part)
        else:
            style = el.find(q("p:style"))
            ref = style.find(q("a:fillRef")) if style is not None else None
            if ref is not None and ref.get("idx", "0") != "0":
                s.fill_kind, s.fill = "solid", env.ctx.resolve(ref)

    def _text(self, s: RShape, txb, ph_type: str | None, inh: list, env: "_Env") -> None:
        chain = env.style_chain(txb, ph_type, inh)
        body = txb.find(q("a:bodyPr"))
        bodies = [body] + [b.find(q("p:txBody")).find(q("a:bodyPr")) for b in inh if b.find(q("p:txBody")) is not None]
        bodies = [b for b in bodies if b is not None]
        s.autofit = "none"
        scale = 1.0
        for b in bodies:
            if b.find(q("a:spAutoFit")) is not None:
                s.autofit = "shape"
                break
            na = b.find(q("a:normAutofit"))
            if na is not None:
                s.autofit = "norm"
                scale = int(na.get("fontScale", 100000)) / 100000   # PowerPoint показывает текст уменьшенным
                break
            if b.find(q("a:noAutofit")) is not None:
                break
        ins = []
        for i, k in enumerate(("lIns", "tIns", "rIns", "bIns")):
            v = next((b.get(k) for b in bodies if b.get(k) is not None), None)
            ins.append(int(v) if v is not None else _DEFAULT_INS[i])
        s.insets = (ins[0] / self.W, ins[1] / self.H, ins[2] / self.W, ins[3] / self.H)
        s.anchor = next((b.get("anchor") for b in bodies if b.get("anchor")), "t")
        s.wrap = next((b.get("wrap") for b in bodies if b.get("wrap")), "square") != "none"
        title_like = ph_type in ("title", "ctrTitle")
        style = s.el.find(q("p:style")) if s.el is not None else None
        font_ref = style.find(q("a:fontRef")) if style is not None else None
        style_color = env.ctx.resolve(font_ref) if font_ref is not None else None   # цвет текста из стиля фигуры
        for p in txb.findall(q("a:p")):
            ppr = p.find(q("a:pPr"))
            lvl = int(ppr.get("lvl", 0)) if ppr is not None else 0
            lvl_pprs = [ppr] + [c.find(q(f"a:lvl{lvl + 1}pPr")) for c in chain]
            lvl_pprs = [x for x in lvl_pprs if x is not None]
            para = RPara(runs=[], level=lvl, align=_first(x.get("algn") for x in lvl_pprs) or "l",
                         bullet=_bullet(lvl_pprs), line_spacing=_line_spacing(lvl_pprs))
            defaults = [ppr.find(q("a:defRPr")) if ppr is not None else None] + [x.find(q("a:defRPr")) for x in lvl_pprs]
            defaults = [d for d in defaults if d is not None]
            for r in p:
                tag = etree.QName(r).localname
                if tag == "br":
                    para.runs.append(self._run(r.find(q("a:rPr")), "\n", defaults, env, title_like, style_color))
                elif tag in ("r", "fld"):
                    t = r.find(q("a:t"))
                    para.runs.append(self._run(r.find(q("a:rPr")), (t.text or "") if t is not None else "", defaults, env,
                                               title_like, style_color))
            if scale < 0.999:
                for r in para.runs:
                    r.size = round(r.size * scale * 2) / 2
            s.paras.append(para)
        if not any(r.text.strip() for p in s.paras for r in p.runs):
            p0 = txb.find(q("a:p"))                   # пустой плейсхолдер: стиль будущего текста — из endParaRPr и цепочки
            ppr = p0.find(q("a:pPr")) if p0 is not None else None
            lvl_pprs = [x for x in [ppr] + [c.find(q("a:lvl1pPr")) for c in chain] if x is not None]
            defaults = [x.find(q("a:defRPr")) for x in lvl_pprs if x.find(q("a:defRPr")) is not None]
            end = p0.find(q("a:endParaRPr")) if p0 is not None else None
            s.default_run = self._run(end, " ", defaults, env, title_like, style_color)

    def _run(self, rpr, text: str, defaults: list, env: "_Env", title_like: bool, style_color: str | None = None) -> RRun:
        cands = ([rpr] if rpr is not None else []) + defaults
        size = _first(c.get("sz") for c in cands)
        bold = _first(c.get("b") for c in cands)
        ital = _first(c.get("i") for c in cands)
        fam = None
        for c in cands:
            lat = c.find(q("a:latin"))
            if lat is not None and lat.get("typeface"):
                fam = lat.get("typeface")
                break
        th = env.ctx.theme
        if fam in (None, "+mj-lt", "+mj-ea"):
            fam = th.major_font if (fam and fam.startswith("+mj")) or (fam is None and title_like) else th.minor_font
        elif fam.startswith("+mn"):
            fam = th.minor_font
        color = None
        for c in cands:
            sf = c.find(q("a:solidFill"))
            if sf is not None:
                color = env.ctx.resolve(sf)
                if color:
                    break
        if color is None:
            color = style_color or env.ctx.scheme("tx1")
        return RRun(text=text, size=int(size) / 100 if size else 18.0, family=fam or "Arial",
                    bold=bold in ("1", "true"), italic=ital in ("1", "true"), color=color)

    def _table(self, s: RShape, gd, env: "_Env") -> None:
        tbl = gd.find(q("a:tbl"))
        if tbl is None:
            return
        rows = []
        font: RRun | None = None
        for tr in tbl.findall(q("a:tr")):
            row = []
            for tc in tr.findall(q("a:tc")):
                texts = ["".join(t.text or "" for t in p.iter(q("a:t"))) for p in tc.iter(q("a:p"))]
                row.append("\n".join(t for t in texts if t))
                if font is None:
                    rpr = tc.find(".//" + q("a:rPr"))
                    if rpr is not None:
                        font = self._run(rpr, "x", [], env, False)
            rows.append(row)
        s.table, s.table_font = rows, font


class _Env:
    """Контекст слайда: цвета, наследуемые плейсхолдеры и цепочка стилей текста."""

    def __init__(self, reader: PresentationReader, ctx: ColorContext, part, layout, master):
        self.reader, self.ctx, self.part = reader, ctx, part
        self.layout, self.master = layout, master
        self._lay = _ph_index(layout._element)
        self._mas = _ph_index(master._element)
        self._tx = master._element.find(q("p:txStyles"))

    def inherited(self, ph_type: str | None, ph_idx: str | None) -> list:
        out = []
        by_idx, by_type = self._lay
        if ph_idx is not None and ph_idx in by_idx:
            out.append(by_idx[ph_idx])
        elif (ph_type or "body") in by_type:
            out.append(by_type[ph_type or "body"])
        mtype = "title" if ph_type in ("title", "ctrTitle") else ph_type if ph_type in ("dt", "ftr", "sldNum") else "body"
        if mtype in self._mas[1]:
            out.append(self._mas[1][mtype])
        return out

    def style_chain(self, txb, ph_type: str | None, inh: list) -> list:
        chain = [txb.find(q("a:lstStyle"))]
        for base in inh:
            bt = base.find(q("p:txBody"))
            if bt is not None:
                chain.append(bt.find(q("a:lstStyle")))
        if self._tx is not None:
            if ph_type in ("title", "ctrTitle"):
                chain.append(self._tx.find(q("p:titleStyle")))
            elif ph_type and ph_type not in ("dt", "ftr", "sldNum"):
                chain.append(self._tx.find(q("p:bodyStyle")))
            else:
                chain.append(self.reader._defaults)
                chain.append(self._tx.find(q("p:otherStyle")))
        else:
            chain.append(self.reader._defaults)
        return [c for c in chain if c is not None]


def _ph_index(root) -> tuple[dict, dict]:
    by_idx, by_type = {}, {}
    tree = root.find(q("p:cSld")).find(q("p:spTree"))
    for el in tree.iter(q("p:sp")):
        ph = el.find(".//" + q("p:ph"))
        if ph is None:
            continue
        if ph.get("idx") is not None:
            by_idx.setdefault(ph.get("idx"), el)
        t = ph.get("type") or "body"
        by_type.setdefault(t, el)
        if t == "ctrTitle":
            by_type.setdefault("title", el)
        if t == "title":
            by_type.setdefault("ctrTitle", el)
    return by_idx, by_type


def _related(part, el, attr: str = q("r:embed")):
    if el is None:
        return None
    rid = el.get(attr)
    if not rid:
        return None
    try:
        return part.related_part(rid)
    except KeyError:
        return None


def _first(values):
    for v in values:
        if v is not None:
            return v
    return None


def _bullet(pprs: list) -> bool:
    for x in pprs:
        if x.find(q("a:buNone")) is not None:
            return False
        if x.find(q("a:buChar")) is not None or x.find(q("a:buAutoNum")) is not None or x.find(q("a:buBlip")) is not None:
            return True
    return False


def _line_spacing(pprs: list) -> float:
    for x in pprs:
        ln = x.find(q("a:lnSpc"))
        if ln is not None:
            pct = ln.find(q("a:spcPct"))
            if pct is not None:
                return int(pct.get("val", 100000)) / 100000
            return 1.0
    return 1.0


def _avg(colors: list[str]) -> str | None:
    if not colors:
        return None
    rs = [int(c[0:2], 16) for c in colors]
    gs = [int(c[2:4], 16) for c in colors]
    bs = [int(c[4:6], 16) for c in colors]
    n = len(colors)
    return f"{sum(rs) // n:02X}{sum(gs) // n:02X}{sum(bs) // n:02X}"
