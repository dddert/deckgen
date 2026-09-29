"""Точка входа слоя парсинга: .pptx -> TemplateModel (DesignSystem + разобранные слайды-образцы).

Кэш — по хэшу файла + версии парсера + режиму классификации (эвристика / VLM с версией промпта).
"""
from __future__ import annotations

import hashlib
import json
from collections import Counter
from pathlib import Path
from typing import Callable

from ..models import SlideKind, TemplateModel, TemplateSlide
from ..ooxml.reader import PresentationReader, RSlide
from .analyzer import analyze_slide
from .design import extract_design, font_usage
from .fonts import TextMeasurer, find_font_file

PARSER_VERSION = "4.4"

Classifier = Callable[[list[TemplateSlide], list[Path]], list[TemplateSlide]]
Renderer = Callable[[Path, Path], list[Path]]


def template_id_of(path: Path) -> str:
    return hashlib.md5(path.read_bytes()).hexdigest()[:10]


_DEMO = {"title": "Заголовок слайда", "ctrTitle": "Название презентации", "subTitle": "Подзаголовок презентации",
         "body": "Первый тезис\nВторой тезис\nТретий тезис", "obj": "Первый тезис\nВторой тезис\nТретий тезис"}


def prepare_template(path: Path, cache_root: Path) -> Path:
    """Шаблон без слайдов-образцов (чистый .pptx/.potx из PowerPoint — только макеты): для каждого макета,
    которого нет среди слайдов, добавляем виртуальный образец с плейсхолдерами и демо-текстом. Геометрия и стиль
    наследуются от макета, дальше парсер/сборщик работают как обычно. Возвращает путь к дополненной копии."""
    from pptx import Presentation
    prs = Presentation(str(path))
    used = {id(s.slide_layout.part) for s in prs.slides}
    missing = [lay for m in prs.slide_masters for lay in m.slide_layouts if id(lay.part) not in used]
    if len(prs.slides) >= 6 or not missing:
        return path
    for lay in missing:
        slide = prs.slides.add_slide(lay)
        for ph in slide.placeholders:
            t = ph.placeholder_format.type
            key = {1: "title", 3: "ctrTitle", 4: "subTitle", 2: "body", 7: "obj"}.get(int(t) if t is not None else 2, None)
            if key and ph.has_text_frame:
                lines = _DEMO[key].split("\n")
                ph.text_frame.text = lines[0]
                for ln in lines[1:]:
                    ph.text_frame.add_paragraph().text = ln
    out = cache_root / template_id_of(path) / "augmented.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


def parse_template(path: str | Path, cache_root: Path, *, use_cache: bool = True, max_shapes: int = 70,
                   renderer: Renderer | None = None, classifier: Classifier | None = None,
                   classifier_key: str = "heur", log: Callable[[str], None] = lambda m: None) -> TemplateModel:
    path = Path(path)
    tid = template_id_of(path)
    source = path
    path = prepare_template(path, cache_root)
    cache = cache_root / tid
    key = f"template_{PARSER_VERSION}_{classifier_key if classifier else 'heur'}.json"
    if use_cache and (cache / key).exists():
        return TemplateModel.model_validate_json((cache / key).read_text(encoding="utf-8"))

    reader = PresentationReader(path)
    rslides = reader.slides()
    log(f"parse: {len(rslides)} слайдов-образцов, {reader.slide_w_in:.2f}×{reader.slide_h_in:.2f} in")

    font_files = {}
    for fam, _ in font_usage(rslides).most_common(4):
        for bold in (False, True):
            p = find_font_file(fam, bold)
            if p:
                font_files[f"{fam}|{'b' if bold else 'r'}"] = str(p)
    measurer = TextMeasurer(font_files)
    logos = _logo_media(rslides)
    photos = _photo_media(reader)
    photo_like = _photo_media(reader, allow_alpha=True)        # только для штрафа образцам с фото в макете
    n_real = _n_slides(source) if path != source else len(rslides)
    tslides = [analyze_slide(rs, measurer, reader.slide_w_in, reader.slide_h_in, logos, max_shapes, photos,
                             reader.read_decor(rs.index), virtual=rs.index >= n_real, photo_like=photo_like,
                             reader=reader)
               for rs in rslides]
    _mark_duplicates(tslides, rslides)

    pngs: list[Path] = []
    if renderer is not None:
        try:
            pngs = renderer(path, cache / "preview")
            for t, p in zip(tslides, pngs):
                t.preview_png = str(p)
        except Exception as e:  # noqa: BLE001 — нет LibreOffice: работаем без превью
            log(f"parse: превью не построены ({e})")
    if classifier is not None and pngs:
        try:
            tslides = classifier(tslides, pngs)
        except Exception as e:  # noqa: BLE001
            log(f"parse: VLM-классификация недоступна ({e}), остаётся эвристика")
    _ensure_minimum(tslides)

    design = extract_design(reader, rslides, tslides, tid, str(path), font_files)
    if path != source:
        design.notes.append(f"В шаблоне не было слайдов-образцов для {len(tslides) - _n_slides(source)} макетов — "
                            "созданы виртуальные образцы по плейсхолдерам макетов.")
    tm = TemplateModel(design=design, slides=tslides, parser_version=PARSER_VERSION)
    cache.mkdir(parents=True, exist_ok=True)
    (cache / key).write_text(tm.model_dump_json(), encoding="utf-8")
    kinds = Counter(t.kind.value for t in tslides if t.usable)
    log(f"parse: типы {dict(kinds)}; шрифты {design.fonts[:3]}; палитра {[c.hex for c in design.palette[:6]]}")
    return tm


def _n_slides(path: Path) -> int:
    from pptx import Presentation
    return len(Presentation(str(path)).slides)


def _logo_media(rslides: list[RSlide]) -> set[str]:
    """Одна и та же картинка в одном месте на 3+ слайдах — логотип/колонтитул."""
    seen: Counter = Counter()
    for sl in rslides:
        for s in sl.shapes:
            if s.kind == "pic" and s.image_part and s.box.area < 0.05:
                seen[(s.image_part, round(s.box.x, 2), round(s.box.y, 2))] += 1
    return {k[0] for k, n in seen.items() if n >= 3}


def _photo_media(reader: PresentationReader, allow_alpha: bool = False) -> set[str]:
    """Непрозрачные фотографии (JPEG, PNG без альфы с фото-подобной палитрой) — демо-контент, а не графика бренда."""
    import io

    from PIL import Image
    out = set()
    for part in reader.prs.part.package.iter_parts():
        name = str(part.partname)
        if not name.startswith("/ppt/media/"):
            continue
        low = name.lower()
        if low.endswith((".jpg", ".jpeg")):
            out.add(name.lstrip("/"))
            continue
        if not low.endswith(".png"):
            continue
        try:
            with Image.open(io.BytesIO(part.blob)) as im:
                alpha = im.mode in ("RGBA", "LA", "P") and (im.mode != "P" or "transparency" in im.info)
                if alpha and not allow_alpha:
                    continue
                if min(im.size) < 300:
                    continue
                small = im.convert("RGBA").resize((64, 64))
                px = [p[:3] for p in (small.get_flattened_data() if hasattr(small, "get_flattened_data") else small.getdata())
                      if p[3] > 200]                       # только непрозрачные пиксели (маска фото не разбавляет подсчёт)
                if len(px) > 400 and len(set(px)) / len(px) > 0.55:   # много разных цветов — фото, а не плоская графика
                    out.add(name.lstrip("/"))
        except Exception:  # noqa: BLE001
            continue
    return out


def _mark_duplicates(tslides: list[TemplateSlide], rslides: list[RSlide]) -> None:
    sig_seen: dict[tuple, int] = {}
    for t, rs in zip(tslides, rslides):
        if not t.usable:
            continue
        sig = (t.kind, t.n_items, t.dark, rs.background.image_part,
               tuple(sorted((s.role.value, round(s.box.x, 2), round(s.box.y, 2), round(s.box.w, 2)) for s in t.texts)),
               tuple(sorted((p.role.value, p.media or "") for p in t.pictures)))
        if sig in sig_seen:
            t.usable, t.reason = False, f"дубликат образца {sig_seen[sig] + 1}"
        else:
            sig_seen[sig] = t.index


def _ensure_minimum(tslides: list[TemplateSlide]) -> None:
    """Колоде нужны хотя бы титул и один контентный слайд — на любом, даже очень бедном шаблоне."""
    usable = [t for t in tslides if t.usable]
    if not usable:
        for t in tslides:
            if t.texts:
                t.usable, t.reason = True, "fallback: других образцов нет"
        usable = [t for t in tslides if t.usable]
    if usable and not any(t.kind == SlideKind.title for t in usable):
        first = min(usable, key=lambda t: t.index)
        first.kind, first.reason = SlideKind.title, "назначен титулом (первый пригодный)"


def summary(tm: TemplateModel) -> dict:
    """Короткое описание шаблона для UI и промптов."""
    d = tm.design
    kinds: dict[str, list[int]] = {}
    for t in tm.usable():
        if t.n_items:
            kinds.setdefault(t.kind.value, []).append(t.n_items)
        else:
            kinds.setdefault(t.kind.value, [])
    return {
        "template_id": d.template_id, "file": Path(d.source_file).name,
        "size_in": [round(d.slide_w_in, 2), round(d.slide_h_in, 2)],
        "fonts": d.fonts[:3], "palette": [{"hex": c.hex, "role": c.role.value} for c in d.palette[:10]],
        "type_scale": [r.model_dump() for r in d.type_scale],
        "slides_total": len(tm.slides), "slides_usable": len(tm.usable()),
        "kinds": {k: sorted(set(v)) for k, v in kinds.items()},
        "notes": d.notes,
    }


def dump_debug(tm: TemplateModel) -> str:
    rows = []
    for t in tm.slides:
        rows.append({"#": t.index + 1, "kind": t.kind.value, "name": t.name, "usable": t.usable, "reason": t.reason,
                     "items": t.n_items, "texts": [(s.role.value, s.item, s.demo_text[:18], s.max_chars) for s in t.texts],
                     "visual": t.visual.source if t.visual else None,
                     "pics": [(p.role.value, p.item) for p in t.pictures]})
    return json.dumps(rows, ensure_ascii=False, indent=1)
