"""Скачать TTF гарнитур, найденных в шаблонах, с Google Fonts (OFL) в assets/fonts — для точной подгонки текста
и отображения в HTML. Встроенные в .pptx шрифты (EOT+MTX) PIL прочитать не может.
Запуск: python scripts/fetch_fonts.py data/templates [Play Inter ...]
"""
from __future__ import annotations

import re
import sys
import urllib.request
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "assets" / "fonts"
API = "https://fonts.googleapis.com/css2?family={}:wght@400;700"


def families_from(folder: Path) -> set[str]:
    from pptx import Presentation
    fams = set()
    for p in folder.glob("*.pptx"):
        for slide in Presentation(str(p)).slides:
            for sh in slide.shapes:
                if sh.has_text_frame:
                    for para in sh.text_frame.paragraphs:
                        fams |= {r.font.name for r in para.runs if r.font.name}
    return {f for f in fams if f and not f.startswith("+")}


def fetch(family: str) -> list[Path]:
    css = urllib.request.urlopen(urllib.request.Request(API.format(family.replace(" ", "+")), headers={
        "User-Agent": "Wget/1.21"}), timeout=20).read().decode()   # старый UA -> Google отдаёт TTF, а не WOFF2
    out = []
    for weight, url in re.findall(r"font-weight:\s*(\d+);.*?src:\s*url\((https://[^)]+\.ttf)\)", css, re.S):
        dst = OUT / f"{family.replace(' ', '')}-{'Bold' if weight == '700' else 'Regular'}.ttf"
        if not dst.exists():
            dst.write_bytes(urllib.request.urlopen(url, timeout=30).read())
        out.append(dst)
    return out


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    args = sys.argv[1:] or ["data/templates"]
    fams = set()
    for a in args:
        fams |= families_from(Path(a)) if Path(a).is_dir() else {a}
    skip = {"Arial", "Calibri", "Times New Roman", "Consolas", "Courier New", "Segoe UI"}
    for f in sorted(fams - skip):
        try:
            got = fetch(f)
            print(f"{f}: {[p.name for p in got] or 'нет на Google Fonts'}")
        except Exception as e:  # noqa: BLE001
            print(f"{f}: не скачан ({e})")
